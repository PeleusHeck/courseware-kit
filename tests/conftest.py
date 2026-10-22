"""测试数据均在 tmp_path 中生成，不使用真实课件。"""

from pathlib import Path
import sys

import pytest
from pypdf import PdfWriter
from pptx import Presentation
from pptx.util import Inches


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


@pytest.fixture
def blank_pdf(tmp_path):
    path = tmp_path / "lecture.pdf"
    with PdfWriter() as writer:
        for index in range(3):
            writer.add_blank_page(width=200 + index, height=300)
        writer.add_outline_item("第一部分", 0)
        writer.add_outline_item("第二部分", 2)
        writer.write(str(path))
    return path


@pytest.fixture
def sample_pptx(tmp_path):
    path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.paragraphs[0].text = "标题"
    paragraph = box.text_frame.add_paragraph()
    paragraph.text, paragraph.level = "要点", 2
    group = slide.shapes.add_group_shape()
    subgroup = group.shapes.add_group_shape()
    subgroup.shapes.add_textbox(
        Inches(1), Inches(2), Inches(4), Inches(1)
    ).text = "组合文本"
    table = slide.shapes.add_table(
        1, 2, Inches(1), Inches(3), Inches(4), Inches(1)
    ).table
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    slide.notes_slide.notes_text_frame.text = "重点复习"
    presentation.save(str(path))
    return path

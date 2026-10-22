"""使用 pypdfium2 和 python-pptx 提取结构化文本，不进行 OCR。"""

from collections import Counter
from collections.abc import Iterable, Iterator
from pathlib import Path
from statistics import median

import pypdfium2 as pdfium
from pypdfium2 import raw as pdfium_raw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.base import BaseShape


def _make_pdfium_line(chars: list[tuple]) -> dict:
    """把字符及字号、坐标转换为后续过滤使用的行和 span。"""
    spans, pending = [], []
    size = chars[0][1]
    for text, font_size, _ in chars:
        if font_size != size:
            spans.append({"text": "".join(pending), "size": size})
            pending = []
            size = font_size
        pending.append(text)
    spans.append({"text": "".join(pending), "size": size})
    boxes = [box for text, _, box in chars if text.strip()]
    if not boxes:
        boxes = [char[2] for char in chars]
    return {"spans": spans, "bbox": (
        min(box[0] for box in boxes), min(box[1] for box in boxes),
        max(box[2] for box in boxes), max(box[3] for box in boxes),
    )}


def _new_text_line(previous: tuple | None, size: float, box: tuple) -> bool:
    if previous is None:
        return False
    old_size, old_box = previous
    tolerance = max(size, old_size, 1)
    return (
        abs((box[1] + box[3] - old_box[1] - old_box[3]) / 2) > tolerance * 0.65
        or box[0] - old_box[2] > tolerance * 3
    )


def _pdfium_lines(textpage, page_box: tuple) -> list[dict]:
    """按字符顺序读取 Unicode；将 PDF 左下原点坐标转为左上原点。"""
    lines, chars = [], []
    previous = None
    left, _, _, top = page_box
    for index in range(textpage.count_chars()):
        code = pdfium_raw.FPDFText_GetUnicode(textpage, index)
        if code in (10, 13):
            if chars:
                lines.append(_make_pdfium_line(chars))
            chars, previous = [], None
            continue
        if not code or (code < 32 and code != 9):
            continue
        text = " " if code == 9 else chr(code)
        if 0xD800 <= code <= 0xDFFF:
            text = "\ufffd"
        size = round(pdfium_raw.FPDFText_GetFontSize(textpage, index), 2)
        x0, y0, x1, y1 = textpage.get_charbox(index)
        box = (x0 - left, top - y1, x1 - left, top - y0)
        if text.strip() and chars and _new_text_line(previous, size, box):
            lines.append(_make_pdfium_line(chars))
            chars, previous = [], None
        chars.append((text, size, box))
        if text.strip():
            previous = (size, box)
    if chars:
        lines.append(_make_pdfium_line(chars))
    return lines


def _group_pdfium_lines(lines: list[dict]) -> list[dict]:
    """按相邻行的字号、间距和缩进重建段落，不跨栏合并。"""
    blocks = []
    previous = None
    for line in lines:
        if not _line_text(line):
            previous = None
            continue
        size = max(span["size"] for span in line["spans"])
        box = line["bbox"]
        same_block = False
        if previous is not None:
            old_size, old_box = previous
            scale = max(size, old_size, 1)
            same_block = (
                abs(size - old_size) <= scale * 0.1
                and -scale * 0.2 <= box[1] - old_box[3] <= scale * 0.8
                and abs(box[0] - old_box[0]) <= scale * 2
            )
        if same_block:
            blocks[-1]["lines"].append(line)
        else:
            blocks.append({"lines": [line]})
        previous = (size, box)
    return blocks


def _read_pdfium_page(document, index: int) -> dict:
    page = document[index]
    try:
        box = page.get_bbox()
        textpage = page.get_textpage()
        try:
            lines = _pdfium_lines(textpage, box)
        finally:
            textpage.close()
        return {"index": index + 1, "height": box[3] - box[1],
                "blocks": _group_pdfium_lines(lines)}
    finally:
        page.close()


def _line_text(line: dict) -> str:
    return "".join(span["text"] for span in line["spans"]).strip()


def _is_margin_line(line: dict, height: float) -> bool:
    """以文本行垂直中心判断它是否位于页面顶部或底部 8%。"""
    center = (line["bbox"][1] + line["bbox"][3]) / 2
    return center < height * 0.08 or center > height * 0.92


def _margin_key(line: dict) -> str:
    return " ".join(_line_text(line).split())


def _repeated_margins(pages: list[dict]) -> set[str]:
    """同一行在同一页出现多次也只计一页，分母包含无文本页面。"""
    counts = Counter()
    for page in pages:
        texts = {
            _margin_key(line)
            for block in page["blocks"]
            for line in block["lines"]
            if _line_text(line) and _is_margin_line(line, page["height"])
        }
        counts.update(texts)
    return {text for text, count in counts.items() if count > len(pages) * 0.6}


def _filtered_pdf_pages(pages: list[dict], margins: set[str]) -> list[dict]:
    filtered = []
    for page in pages:
        blocks = []
        for block in page["blocks"]:
            lines = [
                line for line in block["lines"]
                if _line_text(line) and not (
                    _is_margin_line(line, page["height"])
                    and _margin_key(line) in margins
                )
            ]
            if lines:
                blocks.append(lines)
        filtered.append({"index": page["index"], "blocks": blocks})
    return filtered


def _heading_levels(pages: list[dict]) -> dict[float, int]:
    """以去除页眉页脚后的非空 span 字号中位数估计正文字号。"""
    sizes = [
        span["size"]
        for page in pages for block in page["blocks"] for line in block
        for span in line["spans"]
        if span["text"].strip() and span["size"] > 0
    ]
    if not sizes:
        return {}
    threshold = median(sizes) * 1.2
    headings = sorted({size for size in sizes if size > threshold}, reverse=True)
    return {size: min(index, 3) for index, size in enumerate(headings, start=1)}


def _pdf_blocks(blocks: list[list[dict]], levels: dict[float, int]) -> list[dict]:
    result = []
    for lines in blocks:
        text = "\n".join(_line_text(line) for line in lines).strip()
        if not text or (result and result[-1]["text"] == text):
            continue
        size = max(
            (span["size"] for line in lines for span in line["spans"]
             if span["text"].strip()),
            default=0,
        )
        result.append({"level": levels.get(size, 0), "text": text})
    return result


def extract_pdf(path: str) -> dict:
    """提取文本块；全局映射标题字号，逐页去除连续重复块和空页。"""
    with pdfium.PdfDocument(path) as document:
        raw_pages = [_read_pdfium_page(document, index) for index in range(len(document))]
    filtered = _filtered_pdf_pages(raw_pages, _repeated_margins(raw_pages))
    levels = _heading_levels(filtered)
    pages = []
    for page in filtered:
        blocks = _pdf_blocks(page["blocks"], levels)
        if blocks:
            pages.append({"index": page["index"], "blocks": blocks})
    return {"source": Path(path).name, "type": "pdf", "pages": pages}


def _iter_shapes(shapes: Iterable[BaseShape]) -> Iterator[BaseShape]:
    """按原始顺序深度优先展开组合形状，支持多层嵌套。"""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _table_text(shape: BaseShape) -> str:
    rows = []
    for row in shape.table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def extract_pptx(path: str) -> dict:
    """按形状顺序提取段落和表格，并在每张幻灯片末尾添加备注。"""
    presentation = Presentation(path)
    pages = []
    for index, slide in enumerate(presentation.slides, start=1):
        blocks = []
        for shape in _iter_shapes(slide.shapes):
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        blocks.append({"level": paragraph.level, "text": text})
            if shape.has_table:
                text = _table_text(shape)
                if text:
                    blocks.append({"level": 0, "text": text})
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None and frame.text.strip():
                blocks.append({"level": 0, "text": "备注：" + frame.text.strip()})
        if blocks:
            pages.append({"index": index, "blocks": blocks})
    return {"source": Path(path).name, "type": "pptx", "pages": pages}
